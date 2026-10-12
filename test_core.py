import unittest
from io import BytesIO
from pathlib import Path
from docx import Document
from gift_core import parse_questions, make_gift, escape_gift, read_docx, FormatError


class ConverterTests(unittest.TestCase):
    def test_example_answers_and_feedback(self):
        q = parse_questions(Path(__file__).with_name('ejemplo.txt').read_text(encoding='utf-8'))
        self.assertEqual([x.correct for x in q], ['B', 'A'])
        gift = make_gift(q, 'Parcial/Unidad 1')
        self.assertIn('=Buenos Aires#Buenos Aires es la capital de Argentina.', gift)
        self.assertIn('$CATEGORY: Parcial/Unidad 1\n\n', gift)
        self.assertIn('~Córdoba', gift)
        self.assertEqual(gift.encode('utf-8').decode('utf-8'), gift)

    def test_rejects_incomplete_or_ambiguous_input(self):
        for text in ['¿Pregunta?\nA) Uno\nB) Dos',
                     '¿Pregunta?\nA) Uno\nB) Dos\nOpción correcta: C',
                     '¿Pregunta?\nA) Uno\nA) Dos\nOpción correcta: A',
                     '¿Pregunta?\nA) Uno\nC) Dos\nOpción correcta: A',
                     '¿Pregunta?\nA) Uno\nB) \nOpción correcta: A',
                     '¿Pregunta?\nA) Uno\nB) Dos\nOpción correcta: A y B',
                     '¿Pregunta?\nA) Uno\nB) Dos\ntexto suelto\nOpción correcta: A']:
            with self.subTest(text=text), self.assertRaises(FormatError):
                parse_questions(text)

    def test_final_periods_preserve_answer_and_feedback(self):
        for closing, feedback in [('A.', ''), ('a .', ''),
                                  ('A (Explicación).', 'Explicación'),
                                  ('A. (Explicación).', 'Explicación'),
                                  ('A (Relación (x=2).).', 'Relación (x=2).')]:
            with self.subTest(closing=closing):
                questions = parse_questions('¿Pregunta?\nA) Uno\nB) Dos\nOpción correcta: ' + closing)
                self.assertEqual(questions[0].correct, 'A')
                self.assertEqual(questions[0].feedback, feedback)
                self.assertEqual(make_gift(questions).count('=Uno'), 1)

    def test_ambiguous_answers_are_still_rejected(self):
        for closing in ['A..', 'A y B.', 'A/B', 'A) texto', 'A. texto',
                        'A (Explicación). texto', 'A (sin cierre']:
            with self.subTest(closing=closing), self.assertRaises(FormatError):
                parse_questions('¿Pregunta?\nA) Uno\nB) Dos\nOpción correcta: ' + closing)

    def test_error_numbers_do_not_depend_on_valid_questions(self):
        text = ('1. Primera pregunta\nA) Uno\nB) Dos\nOpción correcta: X\n\n'
                '2. Segunda pregunta\nA) Uno\nB) Dos\nOpción correcta: A y B\n\n'
                '3. Tercera pregunta\nA) Uno\nB) Dos')
        with self.assertRaises(FormatError) as error:
            parse_questions(text)
        message = str(error.exception)
        self.assertIn('Pregunta 1 («Primera pregunta») · línea 4', message)
        self.assertIn('Pregunta 2 («Segunda pregunta») · línea 9', message)
        self.assertIn('Pregunta 3 («Tercera pregunta»): falta el cierre', message)

    def test_special_characters_and_newlines(self):
        self.assertEqual(escape_gift('x={a} ~ b # c: d\\e\nf'), r'x\=\{a\} \~ b \# c\: d\\e\nf')
        q = parse_questions('1. Primera línea\nSegunda línea\nA) x=2\nB) x=3\nRespuesta correcta: a (Relación (x=2))')
        self.assertEqual(q[0].feedback, 'Relación (x=2)')
        self.assertIn(r'[plain]Primera línea\nSegunda línea', make_gift(q))

    def test_docx_preserves_paragraph_and_table_order(self):
        doc = Document()
        doc.add_paragraph('¿Pregunta?')
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = 'A) Sí'
        table.cell(1, 0).text = 'B) No'
        doc.add_paragraph('Opción correcta: A')
        data = BytesIO()
        doc.save(data)
        q = parse_questions(read_docx(data.getvalue()))
        self.assertEqual(q[0].options, {'A': 'Sí', 'B': 'No'})


if __name__ == '__main__':
    unittest.main()
