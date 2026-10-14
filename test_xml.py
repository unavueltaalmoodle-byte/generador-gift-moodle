import base64
from io import BytesIO
from pathlib import Path
import re
import struct
import unittest
import xml.etree.ElementTree as ET
import zlib

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from gift_core import FormatError, parse_questions, make_gift
from image_xml import read_docx_with_images, make_xml


def png(color):
    def chunk(name, body):
        return struct.pack('>I', len(body)) + name + body + struct.pack('>I', zlib.crc32(name + body))
    header = struct.pack('>IIBBBBB', 2, 2, 8, 2, 0, 0, 0)
    pixels = (b'\0' + bytes(color) * 2) * 2
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', header) + chunk(b'IDAT', zlib.compress(pixels)) + chunk(b'IEND', b'')


def save_doc(document):
    data = BytesIO()
    document.save(data)
    return data.getvalue()


class XmlTests(unittest.TestCase):
    def illustrated_document(self):
        doc = Document()
        doc.add_paragraph('1. ¿Qué imagen corresponde? < & >')
        run = doc.add_paragraph().add_run()
        picture = run.add_picture(BytesIO(png((255, 0, 0))), width=Inches(1))
        picture._inline.docPr.set('descr', 'Imagen del enunciado')
        a = doc.add_paragraph('A) Primera ')
        a.add_run().add_picture(BytesIO(png((0, 255, 0))), width=Inches(0.5))
        doc.add_paragraph('B)')
        doc.add_paragraph().add_run().add_picture(BytesIO(png((0, 0, 255))))
        doc.add_paragraph('Opción correcta: B (Explicación con < & >).')
        return doc

    def test_images_stay_in_their_fields_and_base64_is_exact(self):
        content = read_docx_with_images(save_doc(self.illustrated_document()))
        self.assertEqual(len(content.images), 3)
        questions = parse_questions(content.text, content.images)
        self.assertIn('[[IMAGEN_001]]', questions[0].stem)
        self.assertIn('[[IMAGEN_002]]', questions[0].options['A'])
        self.assertIn('[[IMAGEN_003]]', questions[0].options['B'])
        xml = make_xml(questions, 'Prueba/Imágenes', 'Radiología')
        root = ET.fromstring(xml)
        self.assertEqual(root.findtext('question/category/text'), '$course$/Prueba/Imágenes')
        question = root.findall('question')[1]
        self.assertEqual(question.findtext('single'), 'true')
        self.assertEqual(question.findtext('shuffleanswers'), '0')
        self.assertEqual(question.findtext('name/text'), 'Radiología 001')
        self.assertEqual([a.get('fraction') for a in question.findall('answer')], ['0', '100'])
        for container, color in zip([question.find('questiontext'), *question.findall('answer')],
                                    [(255, 0, 0), (0, 255, 0), (0, 0, 255)]):
            files = container.findall('file')
            self.assertEqual(len(files), 1)
            file = files[0]
            self.assertEqual(file.get('path'), '/')
            self.assertEqual(file.get('encoding'), 'base64')
            self.assertEqual(base64.b64decode(file.text, validate=True), png(color))
            self.assertIn('@@PLUGINFILE@@/' + file.get('name'), container.findtext('text'))
        self.assertIn('&lt; &amp; &gt;', question.findtext('questiontext/text'))
        self.assertIn('&lt; &amp; &gt;', question.findall('answer')[1].findtext('feedback/text'))
        self.assertIn('<![CDATA[', xml)
        self.assertNotIn('[[IMAGEN_', xml)
        self.assertTrue(xml.startswith('<?xml version="1.0" encoding="UTF-8"?>'))

    def test_gift_does_not_silently_drop_images(self):
        content = read_docx_with_images(save_doc(self.illustrated_document()))
        with self.assertRaisesRegex(FormatError, 'Moodle XML'):
            make_gift(parse_questions(content.text, content.images))

    def test_plain_text_xml_and_special_characters(self):
        text = Path(__file__).with_name('ejemplo.txt').read_text(encoding='utf-8')
        root = ET.fromstring(make_xml(parse_questions(text)))
        self.assertEqual(len(root.findall('question')), 2)
        self.assertFalse(root.findall('.//file'))
        special = '¿x < 3 & y > 2? ]]>\nA) <script>alert(1)</script>\nB) No\nOpción correcta: A'
        root = ET.fromstring(make_xml(parse_questions(special), prefix='Prueba & < >'))
        self.assertIn('&lt;script&gt;', root.findtext('question/answer/text'))
        self.assertEqual(root.findtext('question/name/text'), 'Prueba & < > 001')

    def test_multiple_questions_have_separate_files(self):
        doc = self.illustrated_document()
        doc.add_paragraph('2. Otra pregunta')
        doc.add_paragraph().add_run().add_picture(BytesIO(png((255, 0, 0))))
        doc.add_paragraph('A) Sí')
        doc.add_paragraph('B) No')
        doc.add_paragraph('Opción correcta: A.')
        content = read_docx_with_images(save_doc(doc))
        questions = parse_questions(content.text, content.images)
        self.assertEqual([len(q.images) for q in questions], [3, 1])
        root = ET.fromstring(make_xml(questions))
        self.assertEqual(len(root.findall('question')[1].findall('.//file')), 1)

    def test_table_images_and_merged_cells_are_not_duplicated(self):
        doc = Document()
        doc.add_paragraph('¿Qué opción?')
        table = doc.add_table(rows=2, cols=2)
        cell = table.cell(0, 0).merge(table.cell(0, 1))
        cell.text = 'A) '
        cell.paragraphs[0].add_run().add_picture(BytesIO(png((1, 2, 3))))
        table.cell(1, 0).merge(table.cell(1, 1)).text = 'B) No'
        doc.add_paragraph('Opción correcta: A')
        content = read_docx_with_images(save_doc(doc))
        self.assertEqual(len(content.images), 1)
        self.assertEqual(len(parse_questions(content.text, content.images)), 1)

    def test_floating_images_are_rejected(self):
        doc = self.illustrated_document()
        inline = doc.element.body.xpath('.//wp:inline')[0]
        inline.tag = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor'
        with self.assertRaisesRegex(FormatError, 'flotantes'):
            read_docx_with_images(save_doc(doc))

    def test_cropped_and_rotated_images_are_rejected(self):
        doc = self.illustrated_document()
        fill = doc.element.body.xpath('.//pic:blipFill')[0]
        crop = OxmlElement('a:srcRect')
        crop.set('l', '1000')
        fill.append(crop)
        with self.assertRaisesRegex(FormatError, 'recortada'):
            read_docx_with_images(save_doc(doc))
        doc = self.illustrated_document()
        doc.element.body.xpath('.//a:xfrm')[0].set('rot', '5400000')
        with self.assertRaisesRegex(FormatError, 'girada'):
            read_docx_with_images(save_doc(doc))

    def test_invalid_xml_control_characters_have_a_clear_error(self):
        with self.assertRaisesRegex(FormatError, 'carácter de control'):
            make_xml(parse_questions('¿Pregunta?\x01\nA) Sí\nB) No\nOpción correcta: A'))

    def test_external_images_are_rejected(self):
        from docx.oxml.ns import qn
        doc = self.illustrated_document()
        blip = doc.element.body.xpath('.//a:blip')[0]
        del blip.attrib[qn('r:embed')]
        blip.set(qn('r:link'), 'external-image')
        with self.assertRaisesRegex(FormatError, 'vinculada'):
            read_docx_with_images(save_doc(doc))

    def test_missing_marker_file_is_rejected(self):
        with self.assertRaisesRegex(FormatError, 'faltan los archivos'):
            parse_questions('¿Pregunta? [[IMAGEN_001]]\nA) Sí\nB) No\nOpción correcta: A')

    def test_multiple_images_in_own_option_paragraph(self):
        doc = Document()
        doc.add_paragraph('¿Pregunta?')
        doc.add_paragraph('A)')
        run = doc.add_paragraph().add_run()
        for color in [(1, 2, 3), (4, 5, 6)]:
            run.add_picture(BytesIO(png(color)))
        doc.add_paragraph('B) No')
        doc.add_paragraph('Opción correcta: A')
        content = read_docx_with_images(save_doc(doc))
        question = parse_questions(content.text, content.images)[0]
        self.assertEqual(len(re.findall(r'\[\[IMAGEN_', question.options['A'])), 2)


if __name__ == '__main__':
    unittest.main()
