from io import BytesIO
from pathlib import Path
import unittest
import xml.etree.ElementTree as ET
from docx import Document
from gift_core import parse_questions, make_gift, escape_gift, read_docx, FormatError
from image_xml import read_docx_with_images, make_xml
import test_xml

QUESTION = '1. ¿Pregunta?\nA) Uno\nB) Dos\nC) Tres\nOpción correcta: B'


class OptionFeedbackTests(unittest.TestCase):
    def test_individual_feedback_is_preserved_in_each_gift_and_xml_answer(self):
        source = QUESTION + '\nRetroalimentación A: Revisá A.\nRetroalimentación B: Bien B.\nRetroalimentación C: Revisá C.\nRetroalimentación general: Para todos.'
        q = parse_questions(source)[0]
        self.assertEqual(q.option_feedback, {'A': 'Revisá A.', 'B': 'Bien B.', 'C': 'Revisá C.'})
        gift = make_gift([q])
        for entry in ['~Uno#Revisá A.', '=Dos#Bien B.', '~Tres#Revisá C.', '####Para todos.']:
            self.assertIn(entry, gift)
        question = ET.fromstring(make_xml([q])).find('question')
        self.assertEqual([a.findtext('feedback/text') for a in question.findall('answer')],
                         ['<div>Revisá A.</div>', '<div>Bien B.</div>', '<div>Revisá C.</div>'])
        self.assertEqual([a.get('fraction') for a in question.findall('answer')], ['0', '100', '0'])
        self.assertEqual(question.findtext('generalfeedback/text'), '<div>Para todos.</div>')

    def test_parentheses_can_coexist_with_incorrect_option_feedback(self):
        q = parse_questions(QUESTION + ' (Bien.)\nRetroalimentación A: No A.')[0]
        self.assertEqual([q.feedback_for(letter) for letter in q.options], ['No A.', 'Bien.', ''])
        self.assertIn('=Dos#Bien.', make_gift([q]))
        self.assertIn('~Uno#No A.', make_gift([q]))

    def test_duplicate_unknown_empty_or_orphan_fields_fail(self):
        for source in [
            QUESTION + '\nRetroalimentación D: No existe.',
            QUESTION + '\nRetroalimentación A:',
            QUESTION + '\nRetroalimentación A: Una.\nRetroalimentación A: Otra.',
            QUESTION + ' (Bien.)\nRetroalimentación B: También bien.',
            'Retroalimentación A: Huérfana.\n' + QUESTION,
            QUESTION.replace('Opción correcta: B', 'Retroalimentación A: Antes.\nOpción correcta: B'),
            QUESTION + '\nRetroalimentación A: [[IMAGEN_001]]',
            QUESTION + '\nFeedback para la Opción A\nTexto en otro párrafo.',
            QUESTION + '\nFeedback\nOpción A: Algo.',
        ]:
            with self.subTest(source=source), self.assertRaises(FormatError):
                parse_questions(source)

    def test_aliases_and_bulleted_options(self):
        for label in ['Retroalimentación A', 'retroalimentacion a', 'Feedback A', 'Feedback para la Opción A']:
            q = parse_questions(QUESTION.replace('A) Uno', '•\tA) Uno') + '\n'+label+': Revisá.')[0]
            self.assertEqual(q.feedback_for('A'), 'Revisá.')
            self.assertEqual(q.options['A'], 'Uno')

    def test_special_characters_and_question_isolation(self):
        explanation = r'Revisá {x} ~ = # : \ < & >'
        questions = parse_questions(QUESTION + '\nRetroalimentación A: '+explanation+'\n\n'+QUESTION.replace('1.', '2.'))
        self.assertEqual(questions[1].option_feedback, {})
        self.assertIn('~Uno#'+escape_gift(explanation), make_gift(questions))
        xml=ET.fromstring(make_xml(questions))
        self.assertIn('&lt; &amp; &gt;', xml.findtext('question/answer/feedback/text'))
        self.assertEqual(xml.findall('question')[1].findtext('answer/feedback/text'), '<div></div>')

    def test_docx_txt_and_images(self):
        text=QUESTION+'\nRetroalimentación A: No es A.\nRetroalimentación B: Es B.'
        doc=Document()
        for line in text.splitlines():
            doc.add_paragraph(line)
        data=BytesIO()
        doc.save(data)
        for source in [text.encode('utf-8-sig').decode('utf-8-sig'), read_docx(data.getvalue()), read_docx_with_images(data.getvalue()).text]:
            self.assertEqual(parse_questions(source)[0].feedback_for('A'), 'No es A.')
        doc=test_xml.XmlTests().illustrated_document()
        doc.add_paragraph('Retroalimentación A: Revisá la imagen verde.')
        content=read_docx_with_images(test_xml.save_doc(doc))
        questions=parse_questions(content.text,content.images)
        xml=ET.fromstring(make_xml(questions))
        self.assertEqual(len(xml.findall('.//file')),3)
        self.assertEqual(xml.findtext('question/answer/feedback/text'), '<div>Revisá la imagen verde.</div>')

    def test_real_first_four_questions(self):
        source=Path(__file__).with_name('preguntas_1_a_4_formateadas.txt').read_text(encoding='utf-8')
        questions=parse_questions(source)
        self.assertEqual([q.correct for q in questions], ['A','C','C','A'])
        self.assertEqual([len(q.option_feedback) for q in questions],[2,3,3,3])
        self.assertEqual(len(ET.fromstring(make_xml(questions)).findall('question')),4)
        self.assertIn('~No  hay Inteligencia Artificial',make_gift(questions))


if __name__ == '__main__':
    unittest.main()
