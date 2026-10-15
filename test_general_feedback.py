from io import BytesIO
import unittest
import xml.etree.ElementTree as ET
from docx import Document
from gift_core import parse_questions, make_gift, read_docx, escape_gift, FormatError
from image_xml import read_docx_with_images, make_xml
import test_xml

QUESTION = '1. ¿Pregunta?\nA) Uno\nB) Dos\nOpción correcta: B'


class GeneralFeedbackTests(unittest.TestCase):
    def test_same_field_is_preserved_in_text_txt_and_both_docx_readers(self):
        source = QUESTION + '\nRetroalimentación general: Una explicación para todas las respuestas.'
        doc = Document()
        for line in source.splitlines():
            doc.add_paragraph(line)
        data = BytesIO()
        doc.save(data)
        sources = [source, source.encode('utf-8-sig').decode('utf-8-sig'),
                   read_docx(data.getvalue()), read_docx_with_images(data.getvalue()).text]
        for text in sources:
            with self.subTest(text=text):
                questions = parse_questions(text)
                self.assertEqual(questions[0].general_feedback, 'Una explicación para todas las respuestas.')
                self.assertEqual(questions[0].feedback, '')
                gift = make_gift(questions)
                self.assertIn('\n####Una explicación para todas las respuestas.\n}', gift)
                xml = ET.fromstring(make_xml(questions))
                self.assertEqual(xml.findtext('question/generalfeedback/text'), '<div>Una explicación para todas las respuestas.</div>')
                self.assertEqual(xml.findtext('question/answer/feedback/text'), '<div></div>')

    def test_specific_and_general_feedback_are_independent(self):
        questions = parse_questions(QUESTION + ' (Devolución de B).\nRetroalimentación general: Devolución de la pregunta.')
        self.assertEqual(questions[0].feedback, 'Devolución de B')
        self.assertEqual(questions[0].general_feedback, 'Devolución de la pregunta.')
        gift = make_gift(questions)
        self.assertIn('=Dos#Devolución de B\n####Devolución de la pregunta.', gift)
        question = ET.fromstring(make_xml(questions)).find('question')
        self.assertEqual(question.findall('answer')[1].findtext('feedback/text'), '<div>Devolución de B</div>')
        self.assertEqual(question.findtext('generalfeedback/text'), '<div>Devolución de la pregunta.</div>')

    def test_general_feedback_does_not_leak_into_neighboring_questions(self):
        source = (QUESTION + '\nRetroalimentación general: Sólo primera.\n\n'
                  + QUESTION.replace('1.', '2.') + '\n\n'
                  + QUESTION.replace('1.', '3.') + '\nRetroalimentación general: Sólo tercera.')
        questions = parse_questions(source)
        self.assertEqual([q.general_feedback for q in questions], ['Sólo primera.', '', 'Sólo tercera.'])
        self.assertEqual(make_gift(questions).count('####'), 2)
        self.assertEqual([q.findtext('generalfeedback/text') for q in ET.fromstring(make_xml(questions))],
                         ['<div>Sólo primera.</div>', '<div></div>', '<div>Sólo tercera.</div>'])

    def test_optional_field_remains_empty_for_existing_input(self):
        questions = parse_questions(QUESTION)
        self.assertEqual(questions[0].general_feedback, '')
        self.assertNotIn('####', make_gift(questions))

    def test_spelling_variants_and_special_characters(self):
        feedback = r'Revisá x={a} ~ b # c: d\e < & > y ]]>.'
        for label in ['Retroalimentación general', 'retroalimentacion general', 'RETROALIMENTACIÓN GENERAL']:
            with self.subTest(label=label):
                questions = parse_questions(QUESTION + '\n' + label + ': ' + feedback)
                self.assertEqual(questions[0].general_feedback, feedback)
                self.assertIn('####' + escape_gift(feedback), make_gift(questions))
                self.assertIn('&lt; &amp; &gt;', ET.fromstring(make_xml(questions)).findtext('question/generalfeedback/text'))

    def test_invalid_placement_empty_duplicate_or_missing_colon_is_rejected(self):
        for text in [
            'Retroalimentación general: Huérfana.\n' + QUESTION,
            QUESTION.replace('Opción correcta: B', 'Retroalimentación general: Antes.\nOpción correcta: B'),
            QUESTION + '\nRetroalimentación general:',
            QUESTION + '\nRetroalimentación general: Una.\nRetroalimentación general: Otra.',
            QUESTION + '\nRetroalimentación general Sin dos puntos',
            QUESTION + '\nRetroalimentación general: [[IMAGEN_001]]',
            QUESTION + '\n2. Nueva pregunta\nRetroalimentación general: Ya empezó la siguiente.',
            QUESTION + '\nRetroalimentación general: Una.\nTexto separado sin etiqueta.',
        ]:
            with self.subTest(text=text), self.assertRaises(FormatError):
                parse_questions(text)

    def test_word_images_and_general_feedback_coexist(self):
        doc = test_xml.XmlTests().illustrated_document()
        doc.add_paragraph('Retroalimentación general: Compará los rasgos de todas las imágenes.')
        content = read_docx_with_images(test_xml.save_doc(doc))
        questions = parse_questions(content.text, content.images)
        xml = ET.fromstring(make_xml(questions))
        self.assertEqual(len(xml.findall('.//file')), 3)
        self.assertEqual(xml.findtext('question/generalfeedback/text'), '<div>Compará los rasgos de todas las imágenes.</div>')
        self.assertEqual(questions[0].correct, 'B')


if __name__ == '__main__':
    unittest.main()
